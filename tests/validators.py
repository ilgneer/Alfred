"""Checks that an output file is a *valid* file of the expected format."""
from __future__ import annotations

import subprocess
import zipfile
from pathlib import Path

MARKER_WORDS = ("Alfred", "Conversão")


def _text(path: Path) -> str:
    data = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise AssertionError(f"{path.name} não está em UTF-8")


def check_image(path: Path, fmt: str) -> None:
    from PIL import Image

    if fmt == "heic":
        import pillow_heif

        pillow_heif.register_heif_opener()
    with Image.open(path) as im:
        im.load()
        assert im.width > 10 and im.height > 10
        expected = {"jpg": "JPEG", "png": "PNG", "webp": "WEBP", "tiff": "TIFF", "bmp": "BMP",
                    "heic": "HEIF", "avif": "AVIF", "gif": "GIF"}[fmt]
        assert im.format == expected, f"formato {im.format} != {expected}"


def check_pdf(path: Path, need_text: bool = False) -> None:
    import pymupdf

    with pymupdf.open(str(path)) as doc:
        assert doc.page_count >= 1
        if need_text:
            text = "".join(p.get_text() for p in doc)
            assert "Alfred" in text or "Relat" in text, "texto não preservado no PDF"


def check_media(path: Path, ffmpeg: str, want_video: bool) -> None:
    proc = subprocess.run([ffmpeg, "-hide_banner", "-v", "error", "-i", str(path), "-f", "null", "-"],
                          capture_output=True, text=True, stdin=subprocess.DEVNULL)
    assert proc.returncode == 0, proc.stderr[-400:]
    info = subprocess.run([ffmpeg, "-hide_banner", "-i", str(path)], capture_output=True, text=True,
                          stdin=subprocess.DEVNULL).stderr
    if want_video:
        assert "Video:" in info, "sem faixa de vídeo"
    else:
        assert "Audio:" in info and "Video:" not in info.replace("attached pic", ""), "saída de áudio inválida"


def check_zip_mimetype(path: Path, mimetype: str) -> None:
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        assert z.read("mimetype").decode().strip() == mimetype


def check_ole(path: Path) -> None:
    assert path.read_bytes()[:8] == bytes.fromhex("D0CF11E0A1B11AE1"), "não é um arquivo OLE (Office 97)"


def check_output(path: Path, ext: str, ffmpeg: str | None = None, text_marker: bool = True) -> None:
    assert path.exists(), f"{path} não existe"
    if path.is_dir():
        files = [p for p in path.rglob("*") if p.is_file()]
        assert files, "pasta de saída vazia"
        return
    assert path.stat().st_size > 0, "arquivo vazio"
    if ext in ("jpg", "png", "webp", "tiff", "bmp", "heic", "avif"):
        check_image(path, ext)
    elif ext == "pdf":
        check_pdf(path)
    elif ext in ("mp3", "m4a", "wav", "flac", "ogg", "opus", "aiff", "wma"):
        assert ffmpeg
        check_media(path, ffmpeg, want_video=False)
    elif ext in ("mp4", "mov", "mkv", "webm", "avi", "wmv", "gif"):
        assert ffmpeg
        check_media(path, ffmpeg, want_video=True)
    elif ext == "docx":
        from docx import Document

        Document(str(path))
    elif ext == "pptx":
        from pptx import Presentation

        assert len(Presentation(str(path)).slides) >= 1
    elif ext == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path)
        assert wb.worksheets
    elif ext == "epub":
        check_zip_mimetype(path, "application/epub+zip")
    elif ext in ("odt", "ods", "odp"):
        mt = {"odt": "text", "ods": "spreadsheet", "odp": "presentation"}[ext]
        check_zip_mimetype(path, f"application/vnd.oasis.opendocument.{mt}")
    elif ext in ("doc", "xls", "ppt"):
        check_ole(path)
    elif ext == "rtf":
        assert path.read_bytes().lstrip().startswith(b"{\\rtf")
    elif ext in ("txt", "md", "html", "csv", "tsv", "srt", "vtt"):
        text = _text(path)
        assert text.strip(), "texto vazio"
        if ext == "vtt":
            assert text.startswith("WEBVTT")
        if ext == "srt":
            assert "-->" in text and "," in text.split("-->")[0]
    elif ext == "zip":
        with zipfile.ZipFile(path) as z:
            assert z.testzip() is None and z.namelist()
    else:
        raise AssertionError(f"sem validador para .{ext}")
