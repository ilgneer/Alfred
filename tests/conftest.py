"""Shared fixtures: real sample files for every input format.

Files are generated on the fly (no binary blobs in the repository, except the
RAR samples, which cannot be produced without WinRAR). Names contain accents
and spaces on purpose.
"""
from __future__ import annotations

import gzip
import os
import shutil
import subprocess
import sys
import tarfile
import zipfile
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

# Isolate user data during tests
_TMP_HOME = Path(os.environ.get("PYTEST_ALFRED_HOME", ROOT / ".pytest_alfred_home"))
os.environ.setdefault("ALFRED_CONFIG_DIR", str(_TMP_HOME / "config"))
os.environ.setdefault("ALFRED_DATA_DIR", str(_TMP_HOME / "data"))
if sys.platform != "win32" and not os.environ.get("DISPLAY"):
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from alfred.core import engines as engines_mod  # noqa: E402

MARKER = "Alfred Conversão Ação"
BASE = "Relatório de teste"
FIXTURES = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="session")
def qapp_session():
    from PySide6.QtWidgets import QApplication

    app = QApplication.instance() or QApplication(["alfred-tests"])
    yield app


@pytest.fixture(scope="session")
def engines():
    return engines_mod.get(refresh=True)


class SampleFactory:
    """Creates (and caches) one sample file per extension."""

    def __init__(self, root: Path, engines):
        self.root = root
        self.engines = engines
        self.cache: dict[str, Path] = {}

    # -- public -------------------------------------------------------------
    def get(self, ext: str) -> Path:
        if ext in self.cache and self.cache[ext].exists():
            return self.cache[ext]
        maker = getattr(self, f"make_{ext.replace('.', '_')}", None)
        alias = ALIASES.get(ext)
        if maker is None and alias:
            src = self.get(alias)
            dst = self.path(ext)
            shutil.copy2(src, dst)
            self.cache[ext] = dst
            return dst
        if maker is None:
            pytest.skip(f"sem gerador de amostra para .{ext}")
        dst = self.path(ext)
        maker(dst)
        if not dst.exists():
            pytest.skip(f"não foi possível gerar amostra .{ext}")
        self.cache[ext] = dst
        return dst

    def path(self, ext: str) -> Path:
        folder = self.root / "pasta com espaços e acentuação"
        folder.mkdir(parents=True, exist_ok=True)
        return folder / f"{BASE}.{ext}"

    # -- helpers -------------------------------------------------------------
    def _ffmpeg(self, *args):
        if not self.engines.has("ffmpeg"):
            pytest.skip("FFmpeg ausente")
        exe = self.engines.path("ffmpeg")
        subprocess.run([exe, "-hide_banner", "-loglevel", "error", "-y", *map(str, args)], check=True,
                       stdin=subprocess.DEVNULL)

    def _lo_from(self, src_ext: str, fmt: str, dst: Path):
        if not self.engines.has("libreoffice"):
            pytest.skip("LibreOffice ausente")
        src = self.get(src_ext)
        out = dst.parent / "_lo"
        out.mkdir(exist_ok=True)
        profile = (self.root / "lo-profile").as_uri()
        subprocess.run([self.engines.path("libreoffice"), f"-env:UserInstallation={profile}", "--headless",
                        "--convert-to", fmt, "--outdir", str(out), str(src)], check=True,
                       stdin=subprocess.DEVNULL, capture_output=True, timeout=300)
        produced = next(out.glob(f"{src.stem}.*"))
        shutil.move(produced, dst)

    # -- images -------------------------------------------------------------
    def _image(self):
        from PIL import Image, ImageDraw

        im = Image.new("RGBA", (320, 200), (8, 9, 13, 255))
        d = ImageDraw.Draw(im)
        d.ellipse((40, 30, 180, 170), fill=(212, 175, 55, 255))
        d.rectangle((200, 60, 300, 140), fill=(255, 216, 77, 160))
        return im

    def make_png(self, dst):
        self._image().save(dst)

    def make_jpg(self, dst):
        self._image().convert("RGB").save(dst, quality=90)

    def make_webp(self, dst):
        self._image().save(dst, "WEBP")

    def make_tiff(self, dst):
        im = self._image()
        im.save(dst, "TIFF", save_all=True, append_images=[im.rotate(90)])

    def make_bmp(self, dst):
        self._image().convert("RGB").save(dst, "BMP")

    def make_heic(self, dst):
        if not self.engines.has("heif_enc"):
            pytest.skip("codificador HEIC ausente")
        import pillow_heif

        pillow_heif.register_heif_opener()
        self._image().convert("RGB").save(dst, "HEIF", quality=80)

    def make_avif(self, dst):
        if not self.engines.has("avif_enc"):
            pytest.skip("AVIF ausente")
        self._image().save(dst, "AVIF")

    def make_svg(self, dst):
        dst.write_text(
            '<svg xmlns="http://www.w3.org/2000/svg" width="240" height="160" viewBox="0 0 240 160">'
            '<rect width="240" height="160" fill="#08090D"/>'
            '<circle cx="80" cy="80" r="60" fill="#D4AF37"/>'
            f'<text x="150" y="90" fill="#F5F5F5" font-size="18">{MARKER[:6]}</text></svg>',
            encoding="utf-8")

    # -- audio / video ---------------------------------------------------
    def make_wav(self, dst):
        self._ffmpeg("-f", "lavfi", "-i", "sine=frequency=440:duration=2", "-ac", "2", "-c:a", "pcm_s16le", dst)

    def make_mp3(self, dst):
        self._ffmpeg("-i", self.get("wav"), "-c:a", "libmp3lame", "-q:a", "4", dst)

    def make_m4a(self, dst):
        self._ffmpeg("-i", self.get("wav"), "-c:a", "aac", "-b:a", "128k", dst)

    def make_aac(self, dst):
        self._ffmpeg("-i", self.get("wav"), "-c:a", "aac", "-b:a", "128k", "-f", "adts", dst)

    def make_flac(self, dst):
        self._ffmpeg("-i", self.get("wav"), "-c:a", "flac", dst)

    def make_ogg(self, dst):
        self._ffmpeg("-i", self.get("wav"), "-c:a", "libvorbis", dst)

    def make_opus(self, dst):
        self._ffmpeg("-i", self.get("wav"), "-c:a", "libopus", dst)

    def make_aiff(self, dst):
        self._ffmpeg("-i", self.get("wav"), "-c:a", "pcm_s16be", dst)

    def make_wma(self, dst):
        self._ffmpeg("-i", self.get("wav"), "-c:a", "wmav2", "-f", "asf", dst)

    def make_mp4(self, dst):
        self._ffmpeg("-f", "lavfi", "-i", "testsrc=size=320x240:rate=15:duration=2",
                     "-f", "lavfi", "-i", "sine=frequency=330:duration=2",
                     "-c:v", "libx264", "-preset", "ultrafast", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", dst)

    def make_mov(self, dst):
        self._ffmpeg("-i", self.get("mp4"), "-c", "copy", dst)

    def make_mkv(self, dst):
        self._ffmpeg("-i", self.get("mp4"), "-c", "copy", dst)

    def make_webm(self, dst):
        self._ffmpeg("-f", "lavfi", "-i", "testsrc=size=160x120:rate=10:duration=1",
                     "-f", "lavfi", "-i", "sine=duration=1", "-c:v", "libvpx-vp9", "-deadline", "realtime",
                     "-cpu-used", "8", "-c:a", "libopus", "-shortest", dst)

    def make_avi(self, dst):
        self._ffmpeg("-i", self.get("mp4"), "-c:v", "mpeg4", "-c:a", "libmp3lame", dst)

    def make_wmv(self, dst):
        self._ffmpeg("-i", self.get("mp4"), "-c:v", "wmv2", "-c:a", "wmav2", dst)

    def make_gif(self, dst):
        self._ffmpeg("-f", "lavfi", "-i", "testsrc=size=160x120:rate=8:duration=1", dst)

    # -- pdf ------------------------------------------------------------------
    def make_pdf(self, dst):
        import pymupdf

        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((72, 80), "Relatório Alfred", fontsize=24)
        page.insert_text((72, 120), MARKER, fontsize=11)
        page.insert_text((72, 140), "Segundo parágrafo com acentuação: ç ã é ü.", fontsize=11)
        # a small table with ruling lines (for table detection)
        x0, y0, cw, rh = 72, 180, 120, 22
        rows = [["Produto", "Preço"], ["Café", "10"], ["Chá", "7"]]
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                rect = pymupdf.Rect(x0 + c * cw, y0 + r * rh, x0 + (c + 1) * cw, y0 + (r + 1) * rh)
                page.draw_rect(rect, color=(0, 0, 0), width=0.8)
                page.insert_text((rect.x0 + 4, rect.y1 - 6), value, fontsize=10)
        png = self.get("png")
        page.insert_image(pymupdf.Rect(72, 300, 232, 400), filename=str(png))
        page2 = doc.new_page()
        page2.insert_text((72, 80), "Página 2", fontsize=18)
        page2.insert_text((72, 110), MARKER, fontsize=11)
        doc.save(str(dst))
        doc.close()

    # -- documents ------------------------------------------------------------
    def make_docx(self, dst):
        from docx import Document

        d = Document()
        d.add_heading("Relatório Alfred", level=1)
        d.add_paragraph(MARKER)
        d.add_paragraph("Segundo parágrafo com acentuação: ç ã é ü.")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Produto", "Preço"
        t.cell(1, 0).text, t.cell(1, 1).text = "Café", "10"
        d.add_picture(str(self.get("png")))
        d.save(dst)

    def make_doc(self, dst):
        self._lo_from("docx", "doc", dst)

    def make_odt(self, dst):
        self._lo_from("docx", "odt", dst)

    def make_rtf(self, dst):
        dst.write_text(r"{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0 Arial;}}"
                       r"\f0\fs24 Alfred Convers\'e3o A\'e7\'e3o\par Segundo par\'e1grafo.\par}",
                       encoding="ascii")

    def make_txt(self, dst):
        dst.write_text(f"{MARKER}\nSegunda linha com * asterisco e # cerquilha.\n\nNovo parágrafo.\n",
                       encoding="utf-8")

    def make_md(self, dst):
        dst.write_text(f"# Relatório\n\n{MARKER}\n\n- item um\n- item dois\n\n| A | B |\n|---|---|\n| 1 | 2 |\n",
                       encoding="utf-8")

    def make_html(self, dst):
        dst.write_text(f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>Teste</title></head>"
                       f"<body><h1>Relatório</h1><p>{MARKER}</p><ul><li>um</li></ul></body></html>",
                       encoding="utf-8")

    def make_epub(self, dst):
        if not self.engines.has("pandoc"):
            pytest.skip("Pandoc ausente")
        subprocess.run([self.engines.path("pandoc"), str(self.get("md")), "-o", str(dst),
                        "--metadata", "title=Teste"], check=True, stdin=subprocess.DEVNULL)

    # -- presentations -----------------------------------------------------
    def make_pptx(self, dst):
        from pptx import Presentation

        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Relatório Alfred"
        s.placeholders[1].text = MARKER
        s2 = prs.slides.add_slide(prs.slide_layouts[1])
        s2.shapes.title.text = "Slide 2"
        prs.save(dst)

    def make_ppt(self, dst):
        self._lo_from("pptx", "ppt", dst)

    def make_odp(self, dst):
        self._lo_from("pptx", "odp", dst)

    # -- spreadsheets ------------------------------------------------------
    def make_xlsx(self, dst):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Vendas"
        ws.append(["Produto", "Preço", "Observação"])
        ws.append(["Café", 10.5, MARKER])
        ws.append(["Chá", 7, "ação"])
        ws2 = wb.create_sheet("Estoque")
        ws2.append(["Item", "Qtd"])
        ws2.append(["Açúcar", 3])
        wb.save(dst)

    def make_xls(self, dst):
        self._lo_from("xlsx", "xls", dst)

    def make_ods(self, dst):
        self._lo_from("xlsx", "ods", dst)

    def make_csv(self, dst):
        # Brazilian Excel style: semicolon, decimal comma, Windows-1252
        dst.write_bytes(f"Produto;Preço;Observação\nCafé;10,50;{MARKER}\nChá;7;ação\n".encode("cp1252"))

    def make_tsv(self, dst):
        dst.write_text(f"Produto\tPreço\nCafé\t10.5\n{MARKER}\t1\n", encoding="utf-8")

    # -- subtitles -------------------------------------------------------------
    def make_srt(self, dst):
        dst.write_text("1\n00:00:01,000 --> 00:00:02,500\nOlá, <i>mundo</i>!\n\n"
                       f"2\n00:00:03,000 --> 00:00:04,000\n{MARKER}\n", encoding="utf-8")

    def make_vtt(self, dst):
        dst.write_text("WEBVTT\n\nNOTE comentário\n\n00:01.000 --> 00:02.500 align:start\n<v Ana>Olá, mundo!\n\n"
                       f"cue-2\n00:00:03.000 --> 00:00:04.000\n{MARKER}\n", encoding="utf-8")

    # -- archives -------------------------------------------------------------
    def _payload(self) -> Path:
        d = self.root / "conteúdo"
        if not d.exists():
            (d / "sub pasta").mkdir(parents=True)
            (d / "olá.txt").write_text(MARKER, encoding="utf-8")
            (d / "sub pasta" / "dados.csv").write_text("a,b\n1,2\n", encoding="utf-8")
        return d

    def make_zip(self, dst):
        d = self._payload()
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
            for f in d.rglob("*"):
                z.write(f, f.relative_to(d.parent))

    def _tar(self, dst, mode):
        with tarfile.open(dst, mode) as t:
            t.add(self._payload(), arcname="conteúdo")

    def make_tar(self, dst):
        self._tar(dst, "w")

    def make_tar_gz(self, dst):
        self._tar(dst, "w:gz")

    def make_tar_bz2(self, dst):
        self._tar(dst, "w:bz2")

    def make_tar_xz(self, dst):
        self._tar(dst, "w:xz")

    def make_gz(self, dst):
        with gzip.open(dst, "wb") as fh:
            fh.write(MARKER.encode("utf-8") * 10)

    def make_bz2(self, dst):
        import bz2

        dst.write_bytes(bz2.compress(MARKER.encode("utf-8") * 10))

    def make_xz(self, dst):
        import lzma

        dst.write_bytes(lzma.compress(MARKER.encode("utf-8") * 10))

    def make_rar(self, dst):
        shutil.copy2(FIXTURES / "rar5-subdirs.rar", dst)

    def make_7z(self, dst):
        if not self.engines.has("sevenzip"):
            pytest.skip("7-Zip ausente")
        subprocess.run([self.engines.path("sevenzip"), "a", "-bso0", str(dst), str(self._payload())], check=True,
                       stdin=subprocess.DEVNULL)


ALIASES = {
    "jpeg": "jpg", "jpe": "jpg", "jfif": "jpg", "tif": "tiff", "heif": "heic",
    "aif": "aiff", "oga": "ogg", "m4v": "mp4", "htm": "html", "markdown": "md",
    "xlsm": "xlsx", "tgz": "tar.gz", "tbz2": "tar.bz2", "txz": "tar.xz",
}


@pytest.fixture(scope="session")
def samples(tmp_path_factory, engines, qapp_session):
    return SampleFactory(tmp_path_factory.mktemp("amostras"), engines)
