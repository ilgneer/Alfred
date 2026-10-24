"""Content preservation: the converted file carries the original content."""
from __future__ import annotations

import csv
import zipfile

import pytest

from alfred.core import routes
from alfred.core.context import ConvertContext, Options
from alfred.core.converters import convert_one
from tests.conftest import MARKER


def run(engines, src, target_key, tmp_path, method=None, **opts):
    target = routes.T[target_key]
    plan = routes.pick_plan(src, target_key, engines)
    if method:
        plan = next(p for r in routes.routes_for(routes.F.category(src), routes.F.extension(src), target_key)
                    for p in r.plans if p.method == method)
        if not engines.has_all(plan.requires):
            pytest.skip("motor ausente")
    if plan is None:
        pytest.skip("rota indisponível nesta máquina")
    ctx = ConvertContext(engines, Options(output_mode="custom", custom_output_dir=str(tmp_path), **opts))
    ctx.begin(src)
    convert_one(ctx, src, target, plan)
    return ctx.commit()


def pdf_text(path):
    import pymupdf

    with pymupdf.open(str(path)) as d:
        return "".join(p.get_text() for p in d)


def test_docx_to_pdf_keeps_text(samples, engines, tmp_path):
    (out,) = run(engines, samples.get("docx"), "pdf", tmp_path)
    assert MARKER in pdf_text(out)


def test_docx_to_pdf_without_libreoffice_keeps_text(samples, engines, tmp_path):
    (out,) = run(engines, samples.get("docx"), "pdf", tmp_path, method="pandoc_qtpdf")
    assert MARKER in pdf_text(out)


def test_pdf_to_docx_keeps_text(samples, engines, tmp_path):
    from docx import Document

    (out,) = run(engines, samples.get("pdf"), "docx", tmp_path)
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    tables = Document(str(out)).tables
    text += "".join(c.text for t in tables for r in t.rows for c in r.cells)
    assert "Alfred" in text and "Ação" in text


def test_pdf_to_txt_and_md(samples, engines, tmp_path):
    (txt,) = run(engines, samples.get("pdf"), "txt", tmp_path / "a")
    assert MARKER in txt.read_text(encoding="utf-8")
    outs = run(engines, samples.get("pdf"), "md", tmp_path / "b")
    md = next(o for o in outs if o.suffix == ".md")
    content = md.read_text(encoding="utf-8")
    assert "# Relatório Alfred" in content
    assert "![](" in content  # embedded image extracted
    media = next(o for o in outs if o.is_dir())
    assert any(media.iterdir())


def test_pdf_pages_300dpi(samples, engines, tmp_path):
    from PIL import Image

    target = routes.T["png_pages"]
    plan = routes.pick_plan(samples.get("pdf"), "png_pages", engines)
    ctx = ConvertContext(engines, Options(output_mode="custom", custom_output_dir=str(tmp_path), pdf_dpi=300))
    ctx.begin(samples.get("pdf"))
    convert_one(ctx, samples.get("pdf"), target, plan)
    (folder,) = ctx.commit()
    pages = sorted(folder.iterdir())
    assert len(pages) == 2
    import pymupdf

    with pymupdf.open(str(samples.get("pdf"))) as d:
        width_pt = d[0].rect.width
    with Image.open(pages[0]) as im:
        assert im.width == pytest.approx(width_pt / 72 * 300, abs=2)
        assert round(im.info["dpi"][0]) == 300


def test_pdf_tables_to_xlsx(samples, engines, tmp_path):
    from openpyxl import load_workbook

    (out,) = run(engines, samples.get("pdf"), "xlsx", tmp_path)
    wb = load_workbook(out)
    values = [c for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c]
    assert "Café" in values and "Preço" in values


def test_csv_brazilian_to_xlsx_types_and_accents(samples, engines, tmp_path):
    from openpyxl import load_workbook

    (out,) = run(engines, samples.get("csv"), "xlsx", tmp_path)
    ws = load_workbook(out).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Produto", "Preço", "Observação")
    assert rows[1][0] == "Café" and rows[1][1] == pytest.approx(10.5) and rows[1][2] == MARKER


def test_xlsx_to_csv_one_file_per_sheet(samples, engines, tmp_path):
    outs = run(engines, samples.get("xlsx"), "csv", tmp_path, method="xlsx_csv")
    names = sorted(o.name for o in outs)
    assert names == ["Relatório de teste - Estoque.csv", "Relatório de teste - Vendas.csv"]
    with open(tmp_path / "Relatório de teste - Vendas.csv", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.reader(fh))
    assert rows[1] == ["Café", "10.5", MARKER]


def test_srt_to_vtt_and_back_keeps_timing(samples, engines, tmp_path):
    (vtt,) = run(engines, samples.get("srt"), "vtt", tmp_path / "v")
    text = vtt.read_text(encoding="utf-8")
    assert "00:00:01.000 --> 00:00:02.500" in text and MARKER in text
    (srt,) = run(engines, vtt, "srt", tmp_path / "s")
    back = srt.read_text(encoding="utf-8")
    assert "1\n00:00:01,000 --> 00:00:02,500" in back


def test_vtt_voice_and_settings(samples, engines, tmp_path):
    (srt,) = run(engines, samples.get("vtt"), "srt", tmp_path)
    text = srt.read_text(encoding="utf-8")
    assert "Ana: Olá, mundo!" in text
    assert "align:start" not in text
    assert "00:00:01,000 --> 00:00:02,500" in text


def test_subtitle_txt_strips_tags(samples, engines, tmp_path):
    (txt,) = run(engines, samples.get("srt"), "txt", tmp_path)
    assert txt.read_text(encoding="utf-8") == f"Olá, mundo!\n{MARKER}\n"


def test_md_to_docx_keeps_text(samples, engines, tmp_path):
    from docx import Document

    (out,) = run(engines, samples.get("md"), "docx", tmp_path)
    assert MARKER in "\n".join(p.text for p in Document(str(out)).paragraphs)


def test_txt_to_md_escapes_markup(samples, engines, tmp_path):
    (out,) = run(engines, samples.get("txt"), "md", tmp_path)
    content = out.read_text(encoding="utf-8")
    assert MARKER in content and "\\*" in content  # literal asterisk preserved


def test_image_dimensions_and_alpha(samples, engines, tmp_path):
    from PIL import Image

    (out,) = run(engines, samples.get("png"), "jpg", tmp_path)
    with Image.open(out) as im:
        assert im.size == (320, 200) and im.mode == "RGB"
    (webp,) = run(engines, samples.get("png"), "webp", tmp_path)
    with Image.open(webp) as im:
        assert im.mode == "RGBA"


def test_svg_to_pdf_is_vector(samples, engines, tmp_path):
    import pymupdf

    (out,) = run(engines, samples.get("svg"), "pdf", tmp_path)
    with pymupdf.open(str(out)) as d:
        page = d[0]
        assert page.get_drawings(), "esperado conteúdo vetorial"
        assert page.rect.width == pytest.approx(240 * 72 / 96, abs=1)


def test_tiff_multipage_to_pdf(samples, engines, tmp_path):
    import pymupdf

    (out,) = run(engines, samples.get("tiff"), "pdf", tmp_path)
    with pymupdf.open(str(out)) as d:
        assert d.page_count == 2


def _duration(engines, path):
    from alfred.core.converters.media import probe

    return probe(engines.path("ffmpeg"), path).duration


def test_audio_duration_preserved(samples, engines, tmp_path):
    for target in ("mp3", "flac", "opus", "wma"):
        (out,) = run(engines, samples.get("wav"), target, tmp_path / target)
        assert _duration(engines, out) == pytest.approx(2.0, abs=0.15), target


def test_video_extract_audio_and_remux(samples, engines, tmp_path):
    (mp3,) = run(engines, samples.get("mp4"), "audio_mp3", tmp_path / "a")
    assert _duration(engines, mp3) == pytest.approx(2.0, abs=0.2)
    (mkv,) = run(engines, samples.get("mp4"), "mkv", tmp_path / "b")
    from alfred.core.converters.media import probe

    assert probe(engines.path("ffmpeg"), mkv).video_codecs[0] == "h264"


def test_zip_extract_contents(samples, engines, tmp_path):
    (folder,) = run(engines, samples.get("zip"), "extract", tmp_path)
    assert (folder / "conteúdo" / "olá.txt").read_text(encoding="utf-8") == MARKER


def test_rar_extract_contents(samples, engines, tmp_path):
    (folder,) = run(engines, samples.get("rar"), "extract", tmp_path)
    files = {p.relative_to(folder).as_posix() for p in folder.rglob("*") if p.is_file()}
    assert "sub/with space/long fn.txt" in files


def test_tar_to_zip_repack(samples, engines, tmp_path):
    (out,) = run(engines, samples.get("tar.gz"), "zip", tmp_path)
    with zipfile.ZipFile(out) as z:
        assert "conteúdo/olá.txt" in z.namelist()
