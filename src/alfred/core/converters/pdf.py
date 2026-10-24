"""PDF conversions (PyMuPDF, pdf2docx, python-pptx, openpyxl, Pandoc)."""
from __future__ import annotations

import os
import statistics
import sys
from pathlib import Path

from alfred.core import formats as F
from alfred.core import paths, procs
from alfred.core.context import ConvertContext
from alfred.core.routes import Plan, Target

MAX_PAGE_PIXELS = 120_000_000


def open_pdf(src: Path):
    import pymupdf

    try:
        doc = pymupdf.open(str(src))
    except Exception as exc:  # noqa: BLE001
        raise RuntimeError(f"Não foi possível abrir o PDF: {exc}") from exc
    if doc.needs_pass:
        doc.close()
        raise RuntimeError("Este PDF é protegido por senha. Remova a senha antes de converter.")
    if doc.page_count == 0:
        doc.close()
        raise RuntimeError("O PDF não possui páginas.")
    return doc


def render_pages(ctx: ConvertContext, src: Path, stem: str, fmt: str, label: str = "páginas") -> Path:
    """Every page as an image (default 300 DPI) inside a dedicated folder."""
    import pymupdf

    dpi = ctx.options.pdf_dpi
    folder = ctx.stage_dir(f"{stem} ({label} {fmt.upper()})")
    with open_pdf(src) as doc:
        n = doc.page_count
        width = max(3, len(str(n)))
        for i, page in enumerate(doc):
            ctx.progress(i / n, f"Página {i + 1} de {n}")
            page_dpi = dpi
            w_in, h_in = page.rect.width / 72, page.rect.height / 72
            if w_in * h_in * page_dpi * page_dpi > MAX_PAGE_PIXELS:
                page_dpi = int((MAX_PAGE_PIXELS / (w_in * h_in)) ** 0.5)
                ctx.notes.append(f"Página {i + 1} muito grande: renderizada a {page_dpi} DPI.")
            pix = page.get_pixmap(matrix=pymupdf.Matrix(page_dpi / 72, page_dpi / 72), alpha=False)
            pix.set_dpi(page_dpi, page_dpi)
            out = folder / f"{'slide' if label == 'slides' else 'página'} {str(i + 1).zfill(width)}.{fmt}"
            if fmt == "jpg":
                pix.save(str(out), jpg_quality=ctx.options.jpg_quality)
            else:
                pix.save(str(out))
    ctx.progress(1.0, "Pronto")
    return folder


def extract_text(src: Path) -> str:
    with open_pdf(src) as doc:
        return "\n\n".join(page.get_text("text", sort=True).strip() for page in doc).strip() + "\n"


def to_markdown(ctx: ConvertContext, src: Path, media_dir: Path | None, media_prefix: str) -> str:
    """Heuristic Markdown: headings by font size, paragraphs per block, images."""
    with open_pdf(src) as doc:
        sizes = []
        for page in doc:
            for b in page.get_text("dict")["blocks"]:
                for line in b.get("lines", []):
                    for span in line["spans"]:
                        if span["text"].strip():
                            sizes.append(round(span["size"], 1))
        body = statistics.median(sizes) if sizes else 11
        parts: list[str] = []
        img_n = 0
        for pno, page in enumerate(doc):
            ctx.progress(pno / doc.page_count, f"Página {pno + 1}")
            blocks = sorted(page.get_text("dict")["blocks"], key=lambda b: (round(b["bbox"][1]), b["bbox"][0]))
            for b in blocks:
                if b.get("type") == 1 and media_dir is not None:
                    data = b.get("image")
                    if not data or b["bbox"][2] - b["bbox"][0] < 24:
                        continue
                    img_n += 1
                    ext = b.get("ext", "png")
                    media_dir.mkdir(parents=True, exist_ok=True)
                    name = f"imagem_{img_n:03d}.{ext}"
                    (media_dir / name).write_bytes(data)
                    parts.append(f"![]({media_prefix}/{name})")
                    continue
                lines = []
                max_size = 0.0
                bold = True
                for line in b.get("lines", []):
                    text = "".join(s["text"] for s in line["spans"]).strip()
                    if text:
                        lines.append(text)
                        for s in line["spans"]:
                            if s["text"].strip():
                                max_size = max(max_size, s["size"])
                                bold = bold and bool(s["flags"] & 16)
                if not lines:
                    continue
                text = lines[0]
                for nxt in lines[1:]:
                    text = text[:-1] + nxt if text.endswith(("-", "­")) else f"{text} {nxt}"
                if max_size >= body * 1.6 and len(text) < 160:
                    parts.append(f"# {text}")
                elif max_size >= body * 1.25 and len(text) < 160:
                    parts.append(f"## {text}")
                elif bold and len(text) < 100 and max_size >= body:
                    parts.append(f"### {text}")
                elif text[:2] in ("• ", "● ", "▪ ", "– "):
                    parts.append("- " + text[2:])
                else:
                    parts.append(text)
            parts.append("")
    return "\n\n".join(p for p in parts).strip() + "\n"


def _pdf2docx_subprocess(ctx: ConvertContext, src: Path, dst: Path) -> None:
    """pdf2docx runs in a child process so the user can cancel it."""
    cmd = paths.app_command("--worker", "pdf2docx", str(src), str(dst))
    env = dict(os.environ)
    if not paths.is_frozen():
        src_root = str(Path(__file__).resolve().parents[3])
        env["PYTHONPATH"] = src_root + os.pathsep + env.get("PYTHONPATH", "")
        cmd[0] = sys.executable  # console-less pythonw is not needed for a child
    ctx.progress(None, "Reconstruindo o layout em DOCX")
    try:
        ctx.run(cmd, env=env, timeout=3 * 3600)
    except procs.ProcessError as exc:
        raise RuntimeError(f"pdf2docx não conseguiu converter: {exc}") from exc


def convert(ctx: ConvertContext, src: Path, target: Target, plan: Plan) -> None:
    stem = F.stem(src)
    key = target.key
    m = plan.method
    if m == "render":
        render_pages(ctx, src, stem, "png" if key == "png_pages" else "jpg")
    elif m == "text":
        ctx.progress(None, "Extraindo texto")
        ctx.stage(f"{stem}.txt").write_text(extract_text(src), encoding="utf-8")
    elif m == "markdown":
        media = f"{stem}_midia"
        md = to_markdown(ctx, src, ctx.staging_dir / media, media)
        ctx.stage(f"{stem}.md").write_text(md, encoding="utf-8")
    elif m == "pdf_epub":
        work = ctx.scratch / "epub"
        work.mkdir(exist_ok=True)
        md = to_markdown(ctx, src, work / "midia", "midia")
        md_file = work / "conteudo.md"
        md_file.write_text(md, encoding="utf-8")
        from alfred.core.converters import office

        office.pandoc(ctx, md_file, "md", ctx.stage(f"{stem}.epub"), "epub", stem, cwd=work)
    elif m == "pdf2docx":
        open_pdf(src).close()  # early, friendly errors (password, empty)
        _pdf2docx_subprocess(ctx, src, ctx.stage(f"{stem}.docx"))
    elif m == "pptx":
        _to_pptx(ctx, src, stem)
    elif m == "tables":
        _to_xlsx(ctx, src, stem)
    else:
        raise RuntimeError(f"Método de PDF desconhecido: {m}")
    ctx.progress(1.0, "Pronto")


def _to_pptx(ctx: ConvertContext, src: Path, stem: str) -> None:
    import pymupdf
    from pptx import Presentation
    from pptx.util import Emu

    with open_pdf(src) as doc:
        first = doc[0].rect
        prs = Presentation()
        # 1 pt = 12700 EMU; PowerPoint accepts 1in..56in
        w = min(max(first.width, 72), 56 * 72)
        h = min(max(first.height, 72), 56 * 72)
        prs.slide_width, prs.slide_height = Emu(int(w * 12700)), Emu(int(h * 12700))
        blank = prs.slide_layouts[6]
        tmp = ctx.scratch / "slides"
        tmp.mkdir(exist_ok=True)
        for i, page in enumerate(doc):
            ctx.progress(i / doc.page_count, f"Página {i + 1} de {doc.page_count}")
            pix = page.get_pixmap(matrix=pymupdf.Matrix(200 / 72, 200 / 72), alpha=False)
            img = tmp / f"{i}.png"
            pix.save(str(img))
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(img), 0, 0, prs.slide_width, prs.slide_height)
            text = page.get_text("text", sort=True).strip()
            if text:
                slide.notes_slide.notes_text_frame.text = text[:30000]
        prs.save(ctx.stage(f"{stem}.pptx"))


def _to_xlsx(ctx: ConvertContext, src: Path, stem: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    found = 0
    with open_pdf(src) as doc:
        for pno, page in enumerate(doc):
            ctx.progress(pno / doc.page_count, f"Procurando tabelas na página {pno + 1}")
            try:
                tables = page.find_tables().tables
            except Exception:  # noqa: BLE001
                tables = []
            for tno, table in enumerate(tables, 1):
                rows = table.extract()
                if not rows:
                    continue
                found += 1
                ws = wb.create_sheet(f"Pág {pno + 1} - Tab {tno}"[:31])
                for row in rows:
                    ws.append([(c or "").strip() if isinstance(c, str) else c for c in row])
        if not found:
            ctx.notes.append(f"{src.name}: nenhuma tabela detectada — exportado o texto por linha.")
            for pno, page in enumerate(doc):
                ws = wb.create_sheet(f"Página {pno + 1}"[:31])
                for line in page.get_text("text", sort=True).splitlines():
                    if line.strip():
                        ws.append([line])
    from alfred.core.converters.documents import _autosize

    for ws in wb.worksheets:
        _autosize(ws)
    wb.save(ctx.stage(f"{stem}.xlsx"))


# ---------------------------------------------------------------------------
# worker entry point (child process)
# ---------------------------------------------------------------------------

def worker_pdf2docx(src: str, dst: str) -> int:
    import logging

    logging.getLogger("pdf2docx").setLevel(logging.ERROR)
    from pdf2docx import Converter

    cv = Converter(src)
    try:
        cv.convert(dst, start=0, end=None)
    finally:
        cv.close()
    return 0 if Path(dst).exists() and Path(dst).stat().st_size > 0 else 2
