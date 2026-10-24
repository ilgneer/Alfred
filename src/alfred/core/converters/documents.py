"""Text documents, presentations and spreadsheets."""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
import shutil
from pathlib import Path

from alfred.core import formats as F
from alfred.core.context import ConvertContext
from alfred.core.converters import office
from alfred.core.routes import Plan, Target


def _move_into_stage(ctx: ConvertContext, produced: Path, name: str) -> Path:
    dst = ctx.stage(name)
    shutil.move(str(produced), str(dst))
    return dst


# ---------------------------------------------------------------------------
# Documents
# ---------------------------------------------------------------------------

def convert_document(ctx: ConvertContext, src: Path, target: Target, plan: Plan) -> None:
    ext = F.extension(src)
    stem = F.stem(src)
    fmt = target.ext
    method = plan.method

    if method == "lo":
        produced = office.lo_convert(ctx, src, fmt, infilter=_lo_infilter_text(ext))
        _move_into_stage(ctx, produced[0], f"{stem}.{fmt}")
    elif method == "pandoc":
        _pandoc_to(ctx, src, ext, fmt, stem)
    elif method == "pandoc_lo":  # md/epub -> doc
        mid = ctx.scratch / "intermediario.docx"
        office.pandoc(ctx, src, ext, mid, "docx", stem)
        produced = office.lo_convert(ctx, mid, fmt)
        _move_into_stage(ctx, produced[0], f"{stem}.{fmt}")
    elif method == "lo_pandoc":  # doc -> txt/md/html/epub
        mid = office.lo_convert(ctx, src, "docx")[0]
        _pandoc_to(ctx, mid, "docx", fmt, stem)
    elif method == "pandoc_qtpdf":
        work = ctx.scratch / "html"
        work.mkdir(exist_ok=True)
        html_file = work / "documento.html"
        _pandoc_html_for_pdf(ctx, src, ext, html_file, stem, work)
        office.html_to_pdf(html_file.read_text(encoding="utf-8"), ctx.stage(f"{stem}.pdf"), work, stem)
    elif method == "html_qtpdf":
        office.html_to_pdf(office.read_text(src), ctx.stage(f"{stem}.pdf"), src.parent, stem)
    elif method == "text_qtpdf":
        office.text_to_pdf(office.read_text(src), ctx.stage(f"{stem}.pdf"), stem)
    else:
        raise RuntimeError(f"Método de documento desconhecido: {method}")
    ctx.progress(1.0, "Pronto")


def _lo_infilter_text(ext: str) -> str | None:
    if ext == "txt":
        return "Text (encoded):UTF8"
    return None


def _pandoc_to(ctx: ConvertContext, src: Path, ext: str, fmt: str, stem: str) -> None:
    if fmt == "md":
        # images extracted next to the Markdown file, referenced relatively
        dst = ctx.stage(f"{stem}.md")
        office.pandoc(ctx, src, ext, dst, "md", stem, extract_media=f"{stem}_midia", cwd=ctx.staging_dir)
        return
    office.pandoc(ctx, src, ext, ctx.stage(f"{stem}.{fmt}"), fmt, stem)


def _pandoc_html_for_pdf(ctx, src: Path, ext: str, html_file: Path, stem: str, work: Path) -> None:
    exe = ctx.engines.path("pandoc")
    inp, reader = office.prepare_pandoc_input(ctx, src, ext)
    ctx.progress(None, "Pandoc convertendo")
    ctx.run([exe, inp, "-f", reader, "-t", "html4", "--standalone", "-o", html_file,
             "--extract-media=midia", f"--resource-path={src.parent}",
             "--metadata", f"pagetitle={stem}"], cwd=work, timeout=900)


# ---------------------------------------------------------------------------
# Presentations
# ---------------------------------------------------------------------------

def convert_presentation(ctx: ConvertContext, src: Path, target: Target, plan: Plan) -> None:
    stem = F.stem(src)
    if plan.method == "lo":
        produced = office.lo_convert(ctx, src, target.ext)
        _move_into_stage(ctx, produced[0], f"{stem}.{target.ext}")
    elif plan.method == "lo_render":
        pdf = office.lo_convert(ctx, src, "pdf")[0]
        from alfred.core.converters import pdf as pdfconv

        pdfconv.render_pages(ctx, pdf, stem, "png" if target.key == "png_pages" else "jpg", label="slides")
    else:
        raise RuntimeError(f"Método desconhecido: {plan.method}")
    ctx.progress(1.0, "Pronto")


# ---------------------------------------------------------------------------
# Spreadsheets
# ---------------------------------------------------------------------------

def sniff_delimiter(sample: str, default: str = ",") -> str:
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in (";", ",", "\t", "|")}
        best = max(counts, key=counts.get)
        return best if counts[best] else default


def read_delimited(path: Path) -> tuple[list[list[str]], str]:
    text = office.read_text(path)
    ext = F.extension(path)
    delim = "\t" if ext == "tsv" else sniff_delimiter(text[:20000])
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    return rows, delim


_INT = re.compile(r"^-?(0|[1-9]\d{0,14})$")
_FLOAT_DOT = re.compile(r"^-?\d+\.\d+$")
_FLOAT_BR = re.compile(r"^-?\d{1,3}(\.\d{3})*,\d+$|^-?\d+,\d+$")


def typed_value(cell: str, delim: str):
    s = cell.strip()
    if not s:
        return None
    if _INT.match(s):
        return int(s)
    if delim != ";" and _FLOAT_DOT.match(s):
        return float(s)
    if delim == ";" and _FLOAT_BR.match(s):
        return float(s.replace(".", "").replace(",", "."))
    return cell


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VERDADEIRO" if v else "FALSO"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat(sep=" ") if isinstance(v, dt.datetime) else v.isoformat()
    return str(v)


def _write_delimited(path: Path, rows, delim: str) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=delim, lineterminator="\r\n")
        for row in rows:
            w.writerow([_cell_text(v) for v in row])


def _autosize(ws) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 500)):
        for cell in row:
            if cell.value is not None:
                widths[cell.column] = max(widths.get(cell.column, 0), min(60, len(str(cell.value))))
    from openpyxl.utils import get_column_letter

    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(8, w + 2)


def _sheet_outputs(ctx, stem: str, fmt: str, sheets: list[tuple[str, list]]) -> None:
    delim = "\t" if fmt == "tsv" else ","
    non_empty = [(n, r) for n, r in sheets if any(any(v not in (None, "") for v in row) for row in r)]
    sheets = non_empty or sheets[:1]
    for name, rows in sheets:
        fname = f"{stem}.{fmt}" if len(sheets) == 1 else f"{stem} - {name}.{fmt}"
        _write_delimited(ctx.stage(fname), rows, delim)


def convert_spreadsheet(ctx: ConvertContext, src: Path, target: Target, plan: Plan) -> None:
    stem = F.stem(src)
    fmt = target.ext
    m = plan.method
    ctx.progress(None, "Convertendo planilha")
    if m == "lo":
        infilter = None
        if F.extension(src) in ("csv", "tsv"):
            text = office.read_text(src)
            delim = "\t" if F.extension(src) == "tsv" else sniff_delimiter(text[:20000])
            utf8 = ctx.scratch / f"entrada-utf8.{F.extension(src)}"
            utf8.write_text(text, encoding="utf-8")
            src = utf8
            infilter = f"Text - txt - csv (StarCalc):{ord(delim)},34,76,1"
        produced = office.lo_convert(ctx, src, fmt, infilter=infilter)
        if fmt in ("csv", "tsv"):
            for p in produced:
                sheet = p.stem[len("entrada-"):] if p.stem.startswith("entrada-") else ""
                name = f"{stem}.{fmt}" if len(produced) == 1 or not sheet else f"{stem} - {sheet}.{fmt}"
                _move_into_stage(ctx, p, name)
        else:
            _move_into_stage(ctx, produced[0], f"{stem}.{fmt}")
    elif m == "csv_xlsx":
        from openpyxl import Workbook

        rows, delim = read_delimited(src)
        wb = Workbook()
        ws = wb.active
        ws.title = (stem[:31] or "Planilha").translate(str.maketrans("[]:*?/\\", "_______"))
        for row in rows:
            ws.append([typed_value(c, delim) for c in row])
        _autosize(ws)
        wb.save(ctx.stage(f"{stem}.xlsx"))
    elif m == "xlsx_csv":
        from openpyxl import load_workbook

        wb = load_workbook(src, read_only=True, data_only=True)
        try:
            sheets = [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
        finally:
            wb.close()
        _sheet_outputs(ctx, stem, fmt, sheets)
    elif m in ("xls_csv", "xls_xlsx"):
        import xlrd

        book = xlrd.open_workbook(str(src))
        sheets = []
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                row = []
                for c in range(sh.ncols):
                    cell = sh.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                        except Exception:  # noqa: BLE001
                            row.append(cell.value)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        row.append(bool(cell.value))
                    elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append(None)
                    else:
                        row.append(cell.value)
                rows.append(row)
            sheets.append((sh.name, rows))
        if m == "xls_csv":
            _sheet_outputs(ctx, stem, fmt, sheets)
        else:
            from openpyxl import Workbook

            wb = Workbook()
            wb.remove(wb.active)
            for name, rows in sheets:
                ws = wb.create_sheet((name or "Planilha")[:31])
                for row in rows:
                    ws.append(row)
                _autosize(ws)
            wb.save(ctx.stage(f"{stem}.xlsx"))
    elif m == "delimited":
        rows, _ = read_delimited(src)
        _write_delimited(ctx.stage(f"{stem}.{fmt}"), rows, "\t" if fmt == "tsv" else ",")
    else:
        raise RuntimeError(f"Método de planilha desconhecido: {m}")
    ctx.progress(1.0, "Pronto")
